from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('/home/ubuntu/final-deliverables/AHC_Refreshed_Detailed_SRS_v2.0.docx')
SCREENSHOT = Path('/home/ubuntu/screenshots/webdev-preview-root-1787237667675074768-6628.png')

INK = '0D2931'
OCHRE = 'D98B16'
WARM = 'F8F4EB'
MUTED = '5C6970'
GREEN = '2F6B56'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, INK)
        set_cell_text(cell, header, bold=True, color='FFFFFF', size=8.5)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if row_index % 2 == 1:
                shade(cells[i], WARM)
            set_cell_text(cells[i], value, size=8.2)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(INK if level == 1 else OCHRE)
    return p


def add_body(doc, text, emphasis=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if emphasis and emphasis in text:
        before, after = text.split(emphasis, 1)
        p.add_run(before)
        r = p.add_run(emphasis)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
        p.add_run(after)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(9.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(9.2)
    return p


def add_requirement_table(doc, rows):
    return add_table(doc, ['ID', 'Requirement', 'Acceptance condition'], rows, [0.8, 3.65, 2.25])


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Affordable Housing Cameroon | Detailed SRS v2.0 | August 2026')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.add_run(' | Page ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_footer(section)
    add_page_number(section)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    styles['Normal'].font.size = Pt(9.5)
    styles['Title'].font.name = 'Aptos Display'
    styles['Title'].font.size = Pt(30)
    styles['Title'].font.color.rgb = RGBColor.from_string(INK)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run('AFFORDABLE HOUSING CAMEROON')
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(OCHRE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run('Detailed Software Requirements\nSpecification')
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Product, trust operations, technical design, and launch-readiness baseline')
    r.italic = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    if SCREENSHOT.exists():
        doc.add_paragraph()
        pic = doc.add_picture(str(SCREENSHOT), width=Inches(5.85))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph('Current AHC public marketplace experience — captured from the managed project preview.')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_paragraph()
    add_table(doc, ['Document control', 'Value'], [
        ('Version', '2.0 — refreshed August 2026'),
        ('Prepared for', 'Affordable Housing Cameroon owner and operating team'),
        ('Implementation baseline', 'Managed React/TypeScript, Express/tRPC, Drizzle, MySQL marketplace project'),
        ('Deployment reference', 'affordableho-8aahm5dj.manus.space'),
        ('Current product checkpoint', 'e6551358 — labelled TEST DATA catalogue restoration'),
        ('Classification', 'Operational requirements and design specification; not legal, tenancy, insurance, tax, or financial advice.'),
    ], [2.0, 4.7])
    doc.add_page_break()

    add_heading(doc, '1. Executive Summary')
    add_body(doc, 'Affordable Housing Cameroon (AHC) is a high-trust rental discovery marketplace for Yaoundé and Douala. It is designed around the practical decision a renter must make before spending transport money: whether a home is still available, what cash is required to move in, what can be independently checked, and how to request a safe viewing without exposing the compound door too early.', 'high-trust rental discovery marketplace')
    add_body(doc, 'The current platform places media and Total Move-In Cash Required at the centre of discovery. It combines paid supplier operations, Field Moderator verification, approved-media governance, landmark-only public location, protected viewing requests, safety reporting, bilingual English/French user experience, and operational dashboards. The product does not collect rent, deposits, tenancy money, or landlord–tenant settlement funds.')
    add_table(doc, ['Outcome', 'AHC mechanism', 'Reason for inclusion'], [
        ('Transparent entry cost', 'Itemised cost inputs and a server-validated Total Move-In Cash figure', 'Reduces surprise advance, deposit, agency, and service charges.'),
        ('Fresh public inventory', '14-day reconfirmation lifecycle with protected archival/reconfirmation workflow', 'Reduces wasted travel to stale listings.'),
        ('Evidence-led trust', 'Field Moderator visit, protected evidence, approved public gallery, and explicit provenance', 'Separates a physical check from unsupported availability claims.'),
        ('Media-first decision support', 'Prominent photos/walkthrough, price immediately beneath, facts and amenities', 'Makes discovery faster while retaining cost transparency.'),
        ('Privacy-aware enquiry', 'Landmark-only map, sign-in threshold, safe viewing request, tracked contact intent', 'Avoids exposing a compound door or private staff evidence publicly.'),
    ], [1.35, 2.75, 2.6])

    add_heading(doc, '2. Scope, Objectives, and Boundaries')
    add_body(doc, 'The scope covers public rental discovery; seeker, Agent, Field Moderator, and Admin workspaces; listing and verification lifecycles; service-order payment reconciliation; safety operations; auditable media governance; bilingual interface; notification fallback; local testing support; and operational reporting.')
    add_table(doc, ['In scope', 'Explicit boundary'], [
        ('Rental discovery and filtering', 'Only public, fresh, eligible homes are returned. Exact compound locations are not disclosed publicly.'),
        ('Total Move-In Cash', 'AHC explains the declared amount; it does not guarantee a landlord’s later actions or collect the tenancy funds.'),
        ('Platform service orders', 'AHC processes its own access, credits, verification, and promotion service records—not rent or deposits.'),
        ('Field verification', 'A pass is an operational, evidence-backed visit state; private evidence and documents remain staff-only.'),
        ('Illustrative TEST DATA', 'Demo imagery is explicitly labelled, isolated in a separate provenance path, and cannot qualify a real listing.'),
        ('Local testing', 'A sanitised local snapshot supports testing and excludes production accounts, contacts, payment references, private evidence, and exact locations.'),
    ], [2.1, 4.6])

    add_heading(doc, '3. Stakeholders, Roles, and Access Model')
    add_body(doc, 'Access is enforced on the server through role-aware procedures. Hiding or guessing a URL is not an authorization mechanism. The platform uses the four operational roles below, alongside unauthenticated browsing.')
    add_table(doc, ['Actor', 'Primary activity', 'Key permission', 'Restriction'], [
        ('Anonymous visitor', 'Search public discovery', 'Can browse eligible public cards and filters', 'Cannot open protected detail, report, schedule, or contact through protected flows.'),
        ('Seeker', 'Evaluate and request homes', 'Can use saved favourites, recent viewing history, viewing requests, reports, profile preferences', 'Cannot publish listings or access staff workspaces.'),
        ('Agent', 'Supply and maintain rental inventory', 'Can maintain profile, credits, access, listings, payment references, verification requests, viewing responses', 'Cannot grant verification, reconcile payment, approve payouts, or access other Agents’ records.'),
        ('Field Moderator', 'Perform assigned physical checks', 'Can claim work, capture protected evidence, submit outcome and neighbourhood observations', 'Cannot access payment reconciliation or self-approve role/payout decisions.'),
        ('Admin', 'Govern the service', 'Can manage roles, review listings/reports, reconcile service orders, approve payouts, review operational dashboards', 'Must be provisioned and server-authorised; Admin pages are never public.'),
    ], [1.25, 1.75, 2.45, 1.35])

    add_heading(doc, '4. Public Marketplace Requirements')
    add_requirement_table(doc, [
        ('FR-P01', 'The marketplace shall return only listings that satisfy publication, freshness, and safety eligibility.', 'An ineligible, stale, held, or unapproved listing is absent from public discovery.'),
        ('FR-P02', 'The default discovery view shall not silently impose a budget ceiling.', 'All otherwise eligible homes are shown until a visitor selects a budget filter.'),
        ('FR-P03', 'Cards shall present media first and Total Move-In Cash immediately beneath the media.', 'Media leads each card; the total cash figure is visible before secondary pricing detail.'),
        ('FR-P04', 'Each public listing shall state a substantive property description.', 'A blank, trivial, or omitted description blocks public qualification.'),
        ('FR-P05', 'Each real public listing shall have at least five approved genuine public photos or a published approved walkthrough.', 'A real listing with fewer than five approved photos and no published walkthrough cannot be public.'),
        ('FR-P06', 'Cards and details shall show bedrooms, bathrooms, parking, declared amenities, and Field Moderator neighbourhood observations where present.', 'Facts display without claiming that Agent-declared fields were physically verified.'),
        ('FR-P07', 'The map shall communicate landmark-radius context, not a compound door.', 'No public payload, card, or share preview contains exact private coordinates or address.'),
        ('FR-P08', 'Listing detail shall suggest relevant public alternatives.', 'Suggestions use public city/type/bedroom/price/area context only; no contact, history, or exact location is used.'),
    ])

    add_heading(doc, '5. Media Governance and Listing Quality')
    add_body(doc, 'Media is governed as a trust control rather than a decorative upload. Private Field Moderator evidence, documents, identity material, and internal notes remain outside the public projection. A separate curated public-media collection holds approved gallery images. Publication requires both listing quality and eligibility rules; the legacy photos-count field is not treated as proof of public media quality.')
    add_table(doc, ['Media class', 'Public treatment', 'Qualification effect'], [
        ('Approved genuine public photo', 'May appear in public card/detail gallery after curation.', 'Counts toward the five-photo standard for real listings.'),
        ('Published Field Moderator walkthrough', 'May appear as the public video experience after approval.', 'Can satisfy the media alternative where fewer than five photos exist.'),
        ('Private visit evidence or document', 'Never appears in public discovery.', 'Supports staff review only; never counts as public gallery media by itself.'),
        ('Illustrative TEST DATA image', 'May appear only on explicitly flagged demo listings with visible English/French disclosure.', 'Cannot be called genuine, owner-authorised, or Field Moderator media; it cannot qualify a real listing.'),
    ], [1.85, 2.75, 2.2])
    add_body(doc, 'The Test Data boundary is deliberate. It permits a usable demonstration catalogue without pretending that illustrative images are proof of a real home. Before a public launch, all TEST DATA listings and illustrative media must be removed or replaced with genuine, authorised, approved property media.')

    add_heading(doc, '6. Seeker Experience and Workflow')
    add_table(doc, ['Step', 'User action', 'System behaviour and safeguard'], [
        ('1. Discover', 'Search city, landmark/area wording, budget, property type, bedrooms, and optional filters.', 'Compact default search preserves immediate media browsing; advanced filters remain available on demand.'),
        ('2. Compare', 'Inspect media, Total Move-In Cash, declared facts, neighbourhood information, freshness, and trust passport.', 'The page distinguishes Agent-declared facts, Field Moderator observations, and illustrative TEST DATA media.'),
        ('3. Sign in', 'Open protected property detail or take protected action.', 'Local authentication gives a clear loading state, role-aware feedback, and CAPTCHA enforcement where configured.'),
        ('4. Request viewing', 'Submit a safe viewing request.', 'The Agent receives a privacy-safe schedule request; exact compound details are not exposed through the public map.'),
        ('5. Contact safely', 'Use the recorded lead action when provided.', 'Intent is logged before a WhatsApp handoff; platform does not create a rent payment flow.'),
        ('6. Report concern', 'Submit a safety report after sign-in.', 'Turnstile verification, reporter ownership, status updates, safety-hold rules, and Admin review protect the process.'),
    ], [0.65, 2.2, 3.95])

    add_heading(doc, '7. Agent Operations and Commercial Workflow')
    add_body(doc, 'The Agent workspace is a dedicated protected page. It supports the commercial supplier relationship while preserving separation between commercial entitlement, listing submission, verification, and public publication. A supplier pays AHC for AHC services, not for an automatic public badge or tenancy guarantee.')
    add_table(doc, ['Stage', 'Agent responsibility', 'System/authority control'], [
        ('Account and identity', 'Create profile and upload required business/identity artefacts where requested.', 'Role remains Agent; only Admin may elevate staff roles.'),
        ('Access and credits', 'Choose a plan or welcome entitlement and submit external service-payment reference.', 'No access/credit grant occurs until Admin reconciliation confirms it.'),
        ('Listing submission', 'Provide clear description, declared facts, cost components, landmark context, and media/verification pathway.', 'Server validates cost math and listing-input requirements; public publication is separately gated.'),
        ('Freshness', 'Use reconfirmation to maintain availability.', 'Listings past the freshness period are removed from public discovery until properly reconfirmed.'),
        ('Verification', 'Order individual or route-batch verification for a specific property.', 'A route batch groups nearby appointments operationally; it does not combine evidence, payment, or outcome decisions.'),
        ('Viewing follow-up', 'Respond to protected viewing requests.', 'Agent queue presents factual deadline states such as scheduled, due soon, or overdue.'),
    ], [1.25, 2.8, 2.75])

    add_heading(doc, '8. Field Moderator Workflow')
    add_body(doc, 'Field Moderators are vetted, Admin-provisioned operational contractors. The workflow is designed to ensure that a physical visit produces a reviewable record without turning the Moderator role into a public viewing route or payment administration role.')
    add_requirement_table(doc, [
        ('FR-M01', 'Only an Admin-provisioned Field Moderator may claim or complete a verification assignment.', 'A Seeker or Agent cannot self-select into Moderator access.'),
        ('FR-M02', 'Each property visit shall retain its own outcome, protected evidence, notes, and audit linkage.', 'A route batch never merges one property’s evidence or decision with another.'),
        ('FR-M03', 'Moderator payment share shall remain held until evidence review and required payout approval.', 'Claiming or completing a visit alone cannot trigger a payout.'),
        ('FR-M04', 'Only approved public media may leave the protected evidence boundary.', 'Documents, internal notes, and unapproved assets cannot appear on public cards or share previews.'),
    ])

    add_heading(doc, '9. Admin Governance, Payments, and Receipts')
    add_body(doc, 'The Admin workspace centralises governance that should not sit with Field Moderators. It includes user/role administration, listing and report review, external payment-reference reconciliation for AHC platform services, verification/payout controls, pilot summaries, and readiness guidance. This does not make AHC a landlord, escrow service, or rent collector.')
    add_table(doc, ['Service / decision', 'Current policy baseline', 'Responsible authority'], [
        ('Welcome Bundle', '3,000 XAF; one-time first 30 days; five listing credits; no priority ranking.', 'Agent initiates; Admin confirms external reference.'),
        ('Starter plan', '10,000 XAF; 30 days; five credits; normal ranking.', 'Agent initiates; Admin confirms external reference.'),
        ('Pro plan', '25,000 XAF; 30 days; priority ranking; up to 20 active listings.', 'Agent initiates; Admin confirms external reference.'),
        ('Featured listing', '2,500 XAF for seven days.', 'Admin-governed service order and listing entitlement.'),
        ('Route-batch verification', '5,000 XAF/property; 4,000 held Moderator share; 1,000 AHC share.', 'Admin confirms payment and later payout approval.'),
        ('Individual verification', '7,500 XAF/property; 6,000 held Moderator share; 1,500 AHC share.', 'Admin confirms payment and later payout approval.'),
    ], [1.65, 3.5, 1.65])

    add_heading(doc, '10. Data Model and Lifecycle')
    add_body(doc, 'The managed relational model separates public marketplace data from staff evidence, commercial service records, and user preferences. Data exposure is controlled by server projections rather than by returning raw database rows to the browser.')
    add_table(doc, ['Domain object', 'Core lifecycle', 'Public exposure rule'], [
        ('Listing', 'Draft → review → published → held/archived; reconfirmation sustains freshness.', 'Only a safe public projection; no contact, private notes, private exact location, or protected evidence.'),
        ('Public gallery media', 'Captured/curated → approved → public or withdrawn.', 'Only explicitly approved genuine media is presented for real listings.'),
        ('Illustrative test media', 'Created → linked to explicit TEST DATA fixture → visible with disclosure.', 'Never treated as genuine media and never attached to a real listing.'),
        ('Verification order', 'Pending payment → paid → scheduled → completed/failed/cancelled → audit/payout review.', 'No protected evidence or internal audit detail is public.'),
        ('Agent access and credit', 'Pending → active/past due/suspended/expired; credits issued and consumed.', 'Commercial details are Agent/Admin only.'),
        ('Viewing request', 'Submitted → Agent response → scheduled/declined/expired.', 'A seeker sees own requests; Agent sees assigned request; Admin oversight is role-gated.'),
        ('Safety report', 'Submitted → evaluated → action/notification where applicable.', 'Reporter and investigation details stay protected.'),
    ], [1.55, 3.0, 2.25])

    add_heading(doc, '11. Technical Architecture and Integrations')
    add_body(doc, 'AHC is implemented as a responsive React 19 and TypeScript client using Tailwind CSS and reusable components, served alongside an Express 4 application with tRPC 11 contracts. Drizzle ORM mediates a managed MySQL-compatible database. The server is the authority for role checks, public projections, entitlement state, payment reconciliation, media eligibility, and audit-sensitive operations.')
    add_table(doc, ['Layer', 'Technology / responsibility', 'Key control'], [
        ('Client', 'React 19, TypeScript, Tailwind CSS 4, role-aware route shells, bilingual dictionaries.', 'Never relies on route hiding for authorization; renders only projected data.'),
        ('API and business logic', 'Express 4 and tRPC 11 procedures.', 'Protected procedures enforce role, ownership, status, and input checks.'),
        ('Persistence', 'Drizzle ORM and managed MySQL database.', 'Separate entities for listings, media, payments, verification, reports, notifications, and audit trails.'),
        ('Authentication', 'OAuth/session foundation with local AHC credential flows and JWT/bcrypt controls.', 'Lockout, clear sign-in feedback, role checks, and protected redirection.'),
        ('Maps and sharing', 'Leaflet landmark-radius experience and server-driven OpenGraph preview route.', 'Public map remains approximate; share cards exclude private evidence and exact location.'),
        ('Notifications', 'Resend owner-email fallback; Meta WhatsApp awaits provider approval.', 'Only high-value operational events; documented fallback and delivery audit.'),
        ('Human-verification', 'Cloudflare Turnstile at sensitive sign-in/report paths.', 'Server validates tokens; deployed hostname configuration is required for live challenge rendering.'),
    ], [1.25, 3.35, 2.2])

    add_heading(doc, '12. Security, Privacy, and Non-Functional Requirements')
    add_requirement_table(doc, [
        ('NFR-S01', 'Authorization shall be enforced server-side for all non-public actions.', 'Changing a client route/hash cannot retrieve Admin, Moderator, Agent, or another user’s protected data.'),
        ('NFR-S02', 'The system shall preserve listing-location privacy.', 'Exact compound door and private evidence are absent from all public APIs, maps, share cards, and recommendations.'),
        ('NFR-S03', 'CAPTCHA verification shall fail closed where configured for login/report actions.', 'Missing, invalid, expired, or unconfigured required token is rejected rather than silently accepted.'),
        ('NFR-S04', 'The platform shall be responsive on phone-sized and desktop screens.', 'Search, media, action buttons, labels, cards, and maps remain usable at tested mobile and desktop breakpoints.'),
        ('NFR-S05', 'Bilingual public and operational UI shall use English/French dictionary coverage.', 'Language selection affects supported labels, messages, forms, and disclosures rather than a single isolated section.'),
        ('NFR-S06', 'Public discovery and media governance shall be auditable through automated regression tests.', 'Publication gate, provenance, authorization, and discovery tests pass before a checkpoint is treated as stable.'),
        ('NFR-S07', 'Notifications shall avoid unnecessary private listing or user data.', 'Owner alerts carry only the operational minimum; failed delivery state is recorded without exposing secret credentials.'),
    ])

    add_heading(doc, '13. Quality Assurance and Acceptance Baseline')
    add_body(doc, 'The current project includes a Vitest regression suite and production build validation. The media-first and illustrative TEST DATA update was validated at checkpoint e6551358 with 234 passing tests and one intentional live-provider skip, followed by desktop and mobile marketplace inspection. The test count will evolve with future features; the durable requirement is that a production release should pass its current automated suite, build, and targeted safety/role/media checks.')
    add_table(doc, ['Acceptance area', 'Representative acceptance test'], [
        ('Publication quality', 'A real listing cannot become public without a substantive description and either five approved genuine public photos or a published walkthrough.'),
        ('TEST DATA provenance', 'Illustrative demo media is projected only for explicitly flagged fixtures, is visibly labelled, and cannot be used to qualify a real listing.'),
        ('Cost integrity', 'The computed Total Move-In Cash matches validated cost components and is presented in discovery/detail.'),
        ('Access control', 'Server rejects unauthorised attempts to use Agent, Moderator, Admin, receipt, or protected viewing/report actions.'),
        ('Privacy', 'Public responses omit exact location, private evidence, direct protected contacts, and internal audit notes.'),
        ('Operational readiness', 'Admin indicators reflect genuine operational counts rather than fabricated activity or customer reviews.'),
    ], [1.7, 5.1])

    add_heading(doc, '14. Launch Readiness and Operating Preconditions')
    add_body(doc, 'The software baseline is materially stronger than a simple classified-listing page, but a public launch depends on operational readiness as well as code. The following prerequisites should be completed before advertising the service broadly.')
    add_table(doc, ['Priority', 'Release prerequisite', 'Owner'], [
        ('Critical', 'Remove all TEST DATA listings, test accounts, illustrative media, and non-production payment references or segregate them in an internal environment.', 'Admin / product owner'),
        ('Critical', 'Replace demonstration galleries with genuine owner-authorised and approved Field Moderator media; maintain five-photo or walkthrough rule.', 'Agents, Field Moderators, Admin'),
        ('Critical', 'Configure Cloudflare Turnstile hostname management for the live production domain and perform signed-out/signed-in smoke tests.', 'Admin / technical owner'),
        ('High', 'Set and test official support, escalation, safety-report, reconciliation, refund, and payout operating procedures.', 'Admin / operations lead'),
        ('High', 'Confirm Resend sender/domain configuration and decide whether Meta WhatsApp notification approval will be completed before launch.', 'Admin / technical owner'),
        ('High', 'Perform an owner-led pilot with real Agents, Moderators, and Seekers; record issues through the current pilot dashboards.', 'Product owner'),
        ('Medium', 'Publish domain, privacy, terms, and appropriate consumer/tenancy disclosures after professional review.', 'Product owner / qualified adviser'),
    ], [0.75, 4.95, 1.1])

    add_heading(doc, '15. Recommended Next Increments')
    add_body(doc, 'The next increments should preserve AHC’s distinction between trusted inventory operations and tenancy. Priority should be based on evidence collected during pilot use rather than feature volume alone.')
    add_table(doc, ['Increment', 'Purpose', 'Dependency / safeguard'], [
        ('Complete genuine media intake', 'Build sufficient real, authorised public galleries for each production listing.', 'Media consent, approval workflow, public/private evidence separation.'),
        ('Payment-provider integration', 'Reduce manual service-order reconciliation.', 'Signed callbacks, idempotency, provider onboarding, exception queue, and reconciliation audit.'),
        ('WhatsApp owner alerts', 'Shorten high-value operational response time.', 'Meta approval, consent/opt-out policy, non-sensitive payload design.'),
        ('Pilot analytics review', 'Assess whether people understand Total Move-In Cash and the trust model.', 'Use aggregated, privacy-preserving measures; do not equate raw views with renter demand.'),
        ('Search relevance tuning', 'Improve suggestions and filter ordering from real behaviour.', 'Use only consented, public, or aggregated signals; never weaken privacy boundaries.'),
    ], [1.8, 2.5, 2.5])

    add_heading(doc, '16. Traceability and References')
    add_body(doc, 'This document is an internal product specification based on the current AHC managed-project implementation and its documented operating decisions. It deliberately avoids presenting the platform’s internal test fixtures, current pilot counts, or provider configurations as public market statistics.')
    add_table(doc, ['Reference', 'Purpose'], [
        ('[1] Current AHC managed project checkpoint e6551358', 'Implementation baseline for the isolated illustrative TEST DATA media path, media-first discovery, and current regression state.'),
        ('[2] AHC deployment: https://affordableho-8aahm5dj.manus.space', 'Current managed deployment reference; access and live configuration may differ from the local preview.'),
        ('[3] AHC source artefacts: drizzle/schema.ts, server/db.ts, server/routers.ts, client/src/pages/Home.tsx', 'Primary implementation trace for database boundaries, public projections, procedures, and discovery user experience.'),
        ('[4] AHC Comprehensive SRS and Local Test Guide v1.0', 'Prior operational baseline; superseded where this v2.0 document records current feature decisions.'),
    ], [2.7, 4.1])
    add_body(doc, 'End of specification.')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
