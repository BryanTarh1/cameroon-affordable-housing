from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path('/home/ubuntu/final-deliverables')
OUT_DIR.mkdir(parents=True, exist_ok=True)
ARCH = ROOT / 'docs/diagrams/ahc-system-architecture.png'
FLOW = ROOT / 'docs/diagrams/ahc-local-snapshot-flow.png'
OUT = OUT_DIR / 'AHC_Comprehensive_SRS_and_Local_Test_Guide.docx'

PAPER = 'F7F3E9'
INK = '132C34'
MANGO = 'D78A1D'
MUTED = '5D6D70'
GREEN = '3C7A57'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=INK):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = 'Manrope'
    run.font.size = Pt(8.6)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Affordable Housing Cameroon  |  ')
    run.font.name = 'Manrope'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_margins(section):
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    add_page_number(section.footer.paragraphs[0])


doc = Document()
set_margins(doc.sections[0])
styles = doc.styles
styles['Normal'].font.name = 'Manrope'
styles['Normal'].font.size = Pt(9.2)
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Manrope')
for name in ['Title', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'DM Serif Display' if name != 'Heading 2' else 'Manrope'
    styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), styles[name].font.name)
styles['Title'].font.size = Pt(29)
styles['Title'].font.color.rgb = RGBColor.from_string(INK)
styles['Heading 1'].font.size = Pt(19)
styles['Heading 1'].font.color.rgb = RGBColor.from_string(INK)
styles['Heading 2'].font.size = Pt(11.5)
styles['Heading 2'].font.color.rgb = RGBColor.from_string(MANGO)


def page(title, kicker=None):
    if len(doc.paragraphs) > 1:
        doc.add_page_break()
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(kicker.upper())
        r.bold = True
        r.font.name = 'DM Mono'
        r.font.size = Pt(8.4)
        r.font.color.rgb = RGBColor.from_string(MANGO)
    h = doc.add_heading(title, level=1)
    h.paragraph_format.space_after = Pt(7)


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.04
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.add_run(item)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, INK)
        set_cell_text(c, value, True, 'FFFFFF')
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            shade(cells[i], PAPER if len(t.rows) % 2 else 'FFFFFF')
            set_cell_text(cells[i], value)
    if widths:
        for row in t.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


# 1 — Cover
cover = doc.add_paragraph()
cover.paragraph_format.space_before = Pt(115)
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cover.add_run('AFFORDABLE HOUSING CAMEROON')
r.font.name = 'DM Mono'; r.font.size = Pt(10); r.bold = True; r.font.color.rgb = RGBColor.from_string(MANGO)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Comprehensive Software\nRequirements Specification')
r.font.name = 'DM Serif Display'; r.font.size = Pt(30); r.bold = True; r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Production platform, trust operations, commercial model, and validated local XAMPP testing package')
r.font.name = 'Manrope'; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(MUTED)
doc.add_paragraph()
tbl = table(['Document status', 'Version', 'Prepared for'], [['Implementation baseline', '1.0 — August 2026', 'AHC product owner and operating team']], [2.1, 1.6, 2.8])
para('Classification: Operational product specification. The document describes the currently implemented AHC platform and the approved local testing boundary. It is not legal, financial, payment-provider, or tenancy advice.')

# 2 — Executive summary
page('1. Executive Summary', 'Product context')
para('Affordable Housing Cameroon (AHC) is a high-trust rental marketplace designed for the practical rental conditions of Yaoundé and Douala. It places Total Move-In Cash Required—not only monthly rent—at the centre of discovery, protects exact compound locations, connects genuine leads through tracked WhatsApp links, and applies operational freshness and verification controls to reduce false or stale inventory.')
para('The platform is a paid-from-day-one supplier marketplace. A unified Agent pathway is used for all listing suppliers. Property seekers browse the marketplace before sign-in; Agents, Field Moderators, and Admins enter protected role-specific workspaces using AHC-owned local authentication.')
table(['Outcome', 'AHC mechanism', 'Why it matters'], [
    ['Cost transparency', 'Itemised cost submission and computed move-in cash total', 'Limits surprise advance, deposit, agency, and service demands.'],
    ['Fresh inventory', '14-day reconfirmation lifecycle and automatic archival', 'Reduces wasted seeker travel.'],
    ['Trust signal', 'Evidence-based Field Moderator visit and explicit public labels', 'Separates visited listings from unverified availability claims.'],
    ['Low-friction lead', 'Tracked WhatsApp deep link', 'Matches current local communication practice without in-app chat.'],
])
para('AHC does not hold rent, deposits, or tenancy money. It processes only its own platform-service orders, such as listing access, featured placement, and verification.')

# 3 — Scope
page('2. Scope, Objectives, and Boundaries', 'SRS scope')
para('The scope includes public rental discovery, protected supplier operations, Field Moderator evidence work, Admin governance, local AHC authentication, payment-order reconciliation for platform services, safety reporting, share previews, and a local data-only test package.')
doc.add_heading('2.1 Product objectives', level=2)
bullets([
    'Help seekers compare realistic entry cost before paying taxi fare or committing to a viewing.',
    'Give Agents a structured, paid workflow for publishing, maintaining, and promoting legitimate inventory.',
    'Create an evidence trail for physical verification without exposing private staff proof publicly.',
    'Keep payment governance and payout approval separate from Field Moderator field work.',
    'Enable a safe local test copy that never contains production credentials or production write access.'
])
doc.add_heading('2.2 Explicit exclusions', level=2)
table(['Excluded activity', 'Boundary'], [
    ['Rent, deposits, tenancy settlement', 'Handled directly by rental parties outside AHC.'],
    ['Public staff proof gallery', 'Evidence stays in protected staff workflows.'],
    ['Free supplier tier', 'There is no free publishing tier.'],
    ['Exact public compound address', 'Map display uses approximate landmark context.'],
    ['Direct local-to-cloud database connection', 'Local testing consumes a sanitised HTTP snapshot only.'],
])

# 4 — Stakeholders
page('3. Users, Roles, and Access Model', 'Actors')
para('AHC uses a central local JWT session model backed by bcrypt password hashing. The access model intentionally distinguishes browsing from operational privileges. Role checks are enforced by server procedures rather than hidden links alone.')
table(['Actor', 'Entry point', 'Core permissions', 'Restrictions'], [
    ['Anonymous visitor', 'Public marketplace', 'Search, filters, approximate map, public listing cards', 'Cannot view protected details, submit reports, or contact through tracked lead flow.'],
    ['Seeker', 'Sign-in after property interest', 'Full signed-in detail, Walk-Thru view where available, reports, tracked WhatsApp action', 'Cannot publish or enter staff workspaces.'],
    ['Agent / supplier', 'Protected Agent workspace', 'Buy service orders, manage credits, submit listings, reconfirm, request verification', 'Cannot reconcile payments, grant badges, or access other suppliers’ data.'],
    ['Field Moderator', 'Protected Operations workspace', 'Claim route work, upload proof, pass/fail visits, perform audits', 'No payment reconciliation or automatic payout authority.'],
    ['Admin', 'Protected Admin workspace', 'Governance, role assignment, bans, settings, reconciliation, payout approval, receipts', 'Must be provisioned by the system owner.'],
])
para('Named non-production fixtures support acceptance walkthroughs: Bryan (Admin), Robinson (Field Moderator), Ebot (paid Agent), Tarh (Seeker), Mireille (Agent), Nadege (pending Agent), and Ateh (new Agent applicant). These are demonstration accounts, not production staff records.')

# 5 — Public seeker functionality
page('4. Public Marketplace and Seeker Requirements', 'Functional requirements')
table(['ID', 'Requirement', 'Acceptance condition'], [
    ['FR-S01', 'Search and filter live listings', 'Marketplace filters return only published, fresh inventory matching area and cost intent.'],
    ['FR-S02', 'Show Total Move-In Cash Required', 'Cards display computed total before secondary monthly-rent context.'],
    ['FR-S03', 'Protect location privacy', 'Map markers represent landmark-radius context rather than a compound door.'],
    ['FR-S04', 'Show truthful verification labels', 'Cards state either “Physically verified by AHC” or “Not yet physically verified.”'],
    ['FR-S05', 'Require sign-in at the protected detail threshold', 'Seeker sign-in is requested before sensitive detail, reports, and tracked contact actions.'],
    ['FR-S06', 'Track WhatsApp lead intent', 'A tracked action is recorded before opening the pre-filled WhatsApp link.'],
])
para('The public card also identifies the listing as managed by the responsible Agent and presents a relative freshness phrase such as “Reconfirmed today” or “X days left.” This differentiates availability reconfirmation from an on-site Field Moderator visit.')
doc.add_heading('4.1 Walk-Thru and neighbourhood intelligence', level=2)
para('Signed-in seekers can view an approved Walk-Thru video where one has been supplied and approved. Listing cards can carry operational neighbourhood badges, such as water/power context, road access, junction proximity, and the Zero-Surprise cost seal. These signals describe documented listing facts; they do not guarantee a tenancy outcome.')

# 6 — Agent workflow
page('5. Unified Agent / Supplier Workflow', 'Supplier operations')
para('AHC deliberately uses one supplier pathway. A person who supplies a home—whether acting as a managing Agent or direct supplier—uses the Agent workspace. This avoids a separate Owner interface while preserving moderation, paid access, listing review, freshness, verification, and safety controls.')
table(['Stage', 'Agent action', 'System control'], [
    ['1. Account and profile', 'Registers and signs in through AHC local auth', 'No automatic Moderator/Admin elevation.'],
    ['2. Access purchase', 'Chooses approved plan and submits platform-service reference', 'Order stays pending until Admin confirmation.'],
    ['3. Listing submission', 'Supplies photos, cost components, landmark information, and availability data', 'First publication goes to moderation.'],
    ['4. Publication', 'Maintains own active listings', 'Credit use and Pro 20-active-listing cap are enforced.'],
    ['5. Freshness', 'Uses reconfirmation action', 'After 14 days without reconfirmation, listing is archived.'],
    ['6. Trust upgrade', 'Requests route-batch or individual physical verification', 'Evidence and Admin payout review are required.'],
])
para('The Agent workspace provides read-only freshness date/countdown visibility so suppliers understand when each listing will expire. A global logout control makes deliberate account switching possible during testing and normal operations.')

# 7 — Moderator
page('6. Field Moderator Operations', 'Evidence and audit')
para('Field Moderators are trusted, Admin-provisioned contractors. They are not public applicants and do not receive payment-reconciliation authority. Their role is to conduct on-site checks, capture evidence, compare the physical property to the supplied listing information, and record a pass/fail outcome.')
table(['Control', 'Requirement'], [
    ['Assignment', 'Moderator works from protected route-ready or individual verification assignments.'],
    ['Evidence', 'The visit record contains required proof such as photos, video, notes, and outcome details.'],
    ['Independence', 'Each property has its own evidence, decision, and audit trail—even in a route batch.'],
    ['Audit', 'A separate verifier may conduct a second-check audit on eligible completed visits.'],
    ['Payout hold', 'Moderator compensation is held until evidence review and Admin approval.'],
    ['Privacy', 'Proof media and internal audit detail do not appear on public listing cards.'],
])
para('A seeker cannot self-select into the Moderator role. The operational risk is mitigated by Admin-only provisioning, identity and work vetting, assignment records, evidence requirements, and the ability to suspend or remove a misused account.')

# 8 — Admin
page('7. Admin Governance and Safety Controls', 'Administrative requirements')
para('The Admin workspace is a protected route on the same domain and codebase, but its server operations require an Admin role. Security comes from authorization checks, not from attempting to hide the route. Bryan is the non-production Admin fixture used to validate this boundary.')
table(['Area', 'Admin responsibility'], [
    ['Role governance', 'Provision Moderator and Admin roles; manage bans and account status.'],
    ['Listing governance', 'Review first publication, safety holds, audit status, and policy exceptions.'],
    ['Commercial settings', 'Set governed platform-service prices and the Moderator share percentage.'],
    ['Payment reconciliation', 'Confirm or reject external payment references for AHC platform services only.'],
    ['Payout approval', 'Approve held Moderator share only after evidence and required audit review.'],
    ['Official receipts', 'Access confirmed service-order receipts; Agents access their own eligible receipts.'],
])
para('Seeker reports can trigger safety holds under documented rules. Admin review should consider account age, report patterns, relevant evidence, and possible coordinated false-report behaviour before imposing a lasting supplier sanction.')

# 9 — Commercial
page('8. Commercial Offer and Payment Requirements', 'Paid-from-day-one')
para('The commercial model charges for AHC platform services only. External mobile-money payment references are reconciled by Admin; unconfirmed references grant no access or credits. Field Moderators do not reconcile payments.')
table(['Service', 'Price', 'Entitlement / policy'], [
    ['Welcome Bundle', '3,000 XAF', 'One-time first 30 days; five credits; no priority ranking.'],
    ['Starter', '10,000 XAF', '30 days; five listing credits; normal ranking.'],
    ['Pro', '25,000 XAF', '30 days; priority ranking; maximum 20 active listings.'],
    ['Featured listing', '2,500 XAF', 'Featured placement for seven days.'],
    ['Route-batch verification', '5,000 XAF per property', '4,000 XAF held Moderator share; 1,000 XAF AHC share.'],
    ['Individual verification', '7,500 XAF per property', '6,000 XAF held Moderator share; 1,500 XAF AHC share.'],
])
para('Route-batch pricing is not a single payment for several homes. Each Agent pays for their own property. The batch is an operational grouping of nearby, flexible appointments; it reduces travel duplication but does not combine evidence or verification decisions.')

# 10 — Verification economic logic
page('9. Verification, Payout, and Trust Logic', 'Operational model')
para('The Agent/supplier pays AHC for the verification service because the verification is a commercial trust upgrade for the supplier’s listing. A passed visit may provide the evidence basis for an explicit physical-verification label and approved Walk-Thru presentation, increasing the likelihood of serious, better-informed seeker enquiries.')
table(['Event', 'State', 'Authority'], [
    ['Agent creates verification order', 'Pending payment', 'Agent initiates; no entitlement yet.'],
    ['Admin verifies external reference', 'Paid / scheduled', 'Admin only.'],
    ['Moderator completes visit', 'Passed or failed with protected evidence', 'Assigned Moderator.'],
    ['Audit where applicable', 'Second-verifier review', 'Authorized staff workflow.'],
    ['Payout approval', 'Moderator share approved or held', 'Admin only.'],
])
para('No automatic payment is made to the Moderator simply because a visit was claimed. This protects AHC from payment for incomplete, weak, or misleading evidence. The share remains held until the appropriate review is complete.')
doc.add_heading('9.1 Incentive alignment', level=2)
para('Agents choose verification when faster trust-building, fewer wasted enquiries, a more credible listing, and stronger cost disclosure are commercially valuable. Seekers do not pay the verification fee; they benefit from the resulting signal when deciding whether a viewing merits their travel expense.')

# 11 — Security
page('10. Authentication, Privacy, and Data Protection', 'Security requirements')
table(['Control', 'Implementation requirement'], [
    ['Local authentication', 'AHC-owned JWT cookie session, bcrypt password hashing, 30-day session policy, and failed-login lockout controls.'],
    ['Role boundary', 'Server checks gate Agent, Moderator, and Admin actions.'],
    ['Public privacy', 'Approximate landmark mapping; no public exact compound door.'],
    ['Evidence privacy', 'Field evidence, audit notes, and internal proof remain protected.'],
    ['Payment boundary', 'No rent, deposit, or tenancy money enters AHC payment order flow.'],
    ['Receipt access', 'Only the paying Agent and confirming Admin can retrieve eligible service receipts.'],
    ['Sharing', 'Canonical share URLs use crawlers’ PNG OpenGraph cards without revealing private evidence.'],
])
para('Any payment-provider integration must verify signed callbacks, corroborate transaction status server-side, apply idempotent order grants, retain audit records, and route mismatches into an Admin-only exception queue. A hosted-checkout-first route has been recommended as the initial integration choice, subject to merchant onboarding and commercial terms.')

# 12 — Architecture
page('11. Solution Architecture', 'Technical view')
para('AHC is implemented as a React and TypeScript frontend with Tailwind/shadcn components, a Node.js/Express backend, tRPC procedures, Drizzle ORM, and a managed MySQL/MariaDB database. Leaflet supports landmark-radius mapping. Sharp produces image cards for crawler-compatible OpenGraph share previews.')
if ARCH.exists():
    doc.add_picture(str(ARCH), width=Inches(6.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
else:
    para('Architecture diagram asset was not available during generation; refer to the text description above.')
table(['Layer', 'Primary responsibility'], [
    ['Client', 'Public discovery, responsive cards/map, protected workspace UI, role-aware navigation.'],
    ['Server', 'tRPC contracts, local-auth checks, pricing entitlement logic, protected exports, share route.'],
    ['Database', 'Listings, profiles, credits, orders, verification state, reports, audit events, settings.'],
    ['External boundary', 'WhatsApp lead redirect and future hosted checkout; no in-app rent handling.'],
])

# 13 — Data
page('12. Data Model and Lifecycle', 'Information design')
para('The production data model separates marketplace operations from tenancy transactions. Commercial records relate to AHC services and never represent a landlord–tenant settlement. The local test database is intentionally much smaller and stores only a current sanitised public-listing view.')
table(['Domain object', 'Key lifecycle'], [
    ['Listing', 'Draft → under review → published → archived/held; freshness reconfirmation supports ongoing publication.'],
    ['Agent profile', 'Pending payment / active / past due / suspended / expired; plan tier and expiry control entitlement.'],
    ['Listing credit', 'Issued by confirmed entitlement and consumed by protected listing publication logic.'],
    ['Payment order', 'Created → pending reference → confirmed or rejected; official receipt only when eligible.'],
    ['Verification order', 'Pending payment → paid → scheduled → passed/failed/cancelled; evidence and payout review are protected.'],
    ['Safety report', 'Submitted → evaluated; documented patterns can lead to temporary safety hold and Admin review.'],
])
para('The local snapshot includes only safe presentation fields such as title, city, approximate area, cost figures, bedrooms, verification label, feature/freshness state, selected media URL where approved, and updated timestamp. It excludes accounts, contact details, payment references, evidence, exact locations, and private notes.')

# 14 — Local test package
page('13. Validated Local XAMPP Test Package', 'Local test boundary')
para('A standalone Courtyard Atlas local interface has been prepared for XAMPP. It is intentionally a read-only presentation test, not a replacement for the full Node/React application. It reads the local `ahc_local_test.sanitized_listings` table through a PHP endpoint and shows cards, filters, cost totals, and a safe detail dialog.')
if FLOW.exists():
    doc.add_picture(str(FLOW), width=Inches(6.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
table(['Local component', 'Purpose', 'Safety boundary'], [
    ['Protected server snapshot', 'Maps only approved public listing fields', 'Requires bearer token; rate limited to one request per 60 seconds.'],
    ['PHP importer', 'Downloads snapshot into local MySQL', 'Uses no cloud database connection string.'],
    ['Local MySQL', 'Stores `sanitized_listings`', 'Separate `ahc_local_test` database only.'],
    ['Local test site', 'Renders local cards and detail dialog', 'Cannot create or alter live marketplace records.'],
])
para('The first validated import produced six sanitised listing records in the local database. During local XAMPP maintenance, a documented temporary root-login fallback may be used only for `ahc_local_test`; it should be replaced with the constrained local account after the local MariaDB privilege tables are rebuilt.')

# 15 — Quality
page('14. Non-Functional Requirements and Test Strategy', 'Quality assurance')
table(['Category', 'Requirement / validation'], [
    ['Responsiveness', 'Public and staff flows are reviewed at desktop and constrained mobile widths, including Leaflet touch controls.'],
    ['Availability', 'Freshness controls remove expired inventory from public display rather than showing known-stale listings.'],
    ['Auditability', 'Payment reconciliation, verification evidence, audit, receipt, and safety workflows record accountable state changes.'],
    ['Accessibility', 'Visible logout control, touch-sized actions, clear verification language, and colour/label distinction.'],
    ['Security', 'Protected procedures enforce role checks; local snapshot uses a separate bearer token and contains no private operational records.'],
    ['Regression', 'The implementation baseline has automated unit and router coverage across access, pricing, safety, receipts, shares, and local snapshot behaviour.'],
])
para('Release validation should cover: anonymous search; Seeker sign-in threshold; Agent plan/credit display; listing submission/reconfirmation; Moderator evidence; Admin-only reconciliation; receipt authorization; share preview; mobile layout; and local snapshot import. A successful local import alone does not replace normal production acceptance checks.')

# 16 — Deployment and roadmap
page('15. Deployment, Operations, and Next Steps', 'Operating guide')
doc.add_heading('15.1 Production operating routine', level=2)
bullets([
    'Admin reviews new supplier service orders and confirms only valid external payment references.',
    'Agents reconfirm availability before the 14-day deadline and keep cost information accurate.',
    'Moderators collect complete field evidence; Admin approves any payment share only after review.',
    'Admin investigates safety reports with pattern-aware review rather than automatic permanent penalties.',
    'Team monitors share cards, mobile layout, and protected role boundaries during release checks.'
])
doc.add_heading('15.2 Recommended next implementation gates', level=2)
table(['Priority', 'Next action', 'Decision / dependency'], [
    ['1', 'Integrate hosted checkout', 'Merchant onboarding, signed callbacks, transaction-status verification, and support terms.'],
    ['2', 'Rebuild local XAMPP privilege tables', 'Replace temporary local-root fallback with least-privilege local importer account.'],
    ['3', 'Scale verification routing', 'Define practical travel windows, cancellation/refund policy, and field capacity.'],
    ['4', 'Operational monitoring', 'Track stale listings, report patterns, payment exception rate, and verification turnaround.'],
])
para('This specification should be updated whenever commercial prices, payment-provider terms, role rights, privacy policy, or local testing boundaries change. The implementation must continue to preserve the non-custodial tenancy boundary and protect private evidence from public display.')

# 17 — acceptance
page('16. Acceptance Summary', 'Release baseline')
para('The AHC implementation baseline combines a mobile-oriented discovery experience with operational protections tailored to urban Cameroon rental dynamics. Its value proposition is not merely listing publication: it is transparent move-in cost, freshness, privacy-aware location context, evidence-backed verification, protected role operations, and clear separation between AHC platform services and tenancy money.')
table(['Acceptance area', 'Baseline status'], [
    ['Public discovery and transparent cost emphasis', 'Implemented'],
    ['Local AHC authentication and role separation', 'Implemented'],
    ['Unified paid Agent workflow and plan entitlements', 'Implemented'],
    ['Field verification, evidence, audit, and held payout workflow', 'Implemented'],
    ['Admin-only service-order reconciliation and printable receipt controls', 'Implemented'],
    ['WhatsApp share/lead support and crawler preview cards', 'Implemented'],
    ['Sanitised one-way XAMPP local test snapshot', 'Validated with six local records'],
    ['Direct production database connection from local XAMPP', 'Intentionally excluded'],
])
para('Approval of this document acknowledges the present implementation baseline and its stated operating constraints. It does not authorise uncontrolled access to production data, removal of human evidence review, or use of AHC as a custodian of rent, deposits, or tenancy funds.')

core = doc.core_properties
core.title = 'Affordable Housing Cameroon — Comprehensive SRS and Local Test Guide'
core.subject = 'Product specification and safe local XAMPP testing boundary'
core.author = 'Affordable Housing Cameroon'
core.keywords = 'Affordable Housing Cameroon, SRS, rental marketplace, local testing, XAMPP'
doc.save(OUT)
print(OUT)
