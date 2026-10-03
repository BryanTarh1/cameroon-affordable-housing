from pathlib import Path
from math import ceil
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/cameroon-affordable-housing')
ASSETS = ROOT / 'business-plan-assets'
ASSETS.mkdir(exist_ok=True)

NAVY = '#153243'
TEAL = '#1B7F79'
GOLD = '#D49A3A'
CORAL = '#D96C4F'
SAND = '#F6F0E7'
INK = '#263238'
MUTED = '#66747A'
GREEN = '#3E8E62'

plt.rcParams['font.family'] = 'DejaVu Sans'

def box(ax, x, y, w, h, title, subtitle='', color=NAVY, text_color='white', fontsize=11):
    p = FancyBboxPatch((x,y), w,h, boxstyle='round,pad=0.018,rounding_size=0.025', linewidth=0, facecolor=color)
    ax.add_patch(p)
    ax.text(x+w/2, y+h*0.62, title, ha='center', va='center', color=text_color, fontsize=fontsize, weight='bold')
    if subtitle:
        ax.text(x+w/2, y+h*0.31, subtitle, ha='center', va='center', color=text_color, fontsize=8.5, wrap=True)

def arrow(ax, x1,y1,x2,y2,label='', color=GOLD, rad=0.0):
    a = FancyArrowPatch((x1,y1),(x2,y2), arrowstyle='-|>', mutation_scale=14, linewidth=2, color=color, connectionstyle=f'arc3,rad={rad}')
    ax.add_patch(a)
    if label:
        ax.text((x1+x2)/2, (y1+y2)/2+0.035, label, ha='center', va='center', color=INK, fontsize=8, bbox=dict(facecolor='white', edgecolor='none', alpha=.85, pad=1.5))

def save_fig(fig, name):
    path = ASSETS / name
    fig.savefig(path, dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return path

# 1. Ecosystem flow
fig, ax = plt.subplots(figsize=(12,6.4))
ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off'); fig.patch.set_facecolor('white')
ax.text(.04,.94,'1. Marketplace ecosystem: who creates value and who gets paid', fontsize=17, weight='bold', color=NAVY)
ax.text(.04,.89,'The platform coordinates discovery, verification, viewing, commitment, and settlement.', fontsize=10.5, color=MUTED)
box(ax,.05,.56,.19,.17,'SEEKER','Finds home; pays rent / deposit\nReceives protection + dispute path', TEAL)
box(ax,.30,.56,.19,.17,'OWNER','Provides property\nReceives rent and qualified demand', NAVY)
box(ax,.55,.56,.19,.17,'AGENT','Viewing, negotiation, paperwork\nEarns success-fee share', CORAL)
box(ax,.80,.56,.15,.17,'PLATFORM','Tracks deal, escrow,\nreputation + payouts', GOLD, INK)
box(ax,.17,.20,.21,.16,'MODERATOR','Verifies listing\nEarns upfront fee + deal slice', GREEN)
box(ax,.58,.20,.24,.16,'MOBILE MONEY / PAYMENTS','Moves funds; applies provider\nfees and settlement controls', '#6B7280')
arrow(ax,.24,.64,.30,.64,'demand')
arrow(ax,.49,.64,.55,.64,'qualified lead')
arrow(ax,.74,.64,.80,.64,'recorded deal')
arrow(ax,.19,.56,.27,.36,'verification')
arrow(ax,.80,.56,.70,.36,'payouts')
arrow(ax,.70,.20,.51,.20,'settlement')
ax.text(.04,.07,'Value principle: the platform is not only a listing board; it is a transaction and trust layer.', fontsize=10, color=NAVY, style='italic')
ECO = save_fig(fig,'01_ecosystem_flow.png')

# 2. Commission waterfall
fig, ax = plt.subplots(figsize=(12,5.5))
ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off')
ax.text(.04,.93,'2. Closed-deal commission waterfall', fontsize=17, weight='bold', color=NAVY)
ax.text(.04,.88,'Illustrative assumption: commission equals one month of rent; rent = 40,000 XAF.', fontsize=10.5, color=MUTED)
# left gross
box(ax,.05,.40,.21,.24,'CLOSED DEAL','Commission pool\n40,000 XAF', NAVY, 'white', 13)
# split boxes
box(ax,.38,.58,.20,.22,'AGENT — 70%','28,000 XAF\nViewing + negotiation + paperwork', CORAL, 'white', 11)
box(ax,.38,.30,.20,.22,'PLATFORM — 20%','8,000 XAF\nInfrastructure + trust + support', GOLD, INK, 11)
box(ax,.38,.02,.20,.22,'MODERATOR — 10%','4,000 XAF\nVerified listing + quality signal', GREEN, 'white', 11)
arrow(ax,.26,.52,.38,.69,'split')
arrow(ax,.26,.52,.38,.41,'split')
arrow(ax,.26,.52,.38,.13,'split')
box(ax,.68,.33,.25,.28,'OWNER / SEEKER','Rent and deposit are not\nplatform revenue. They remain\ntransaction funds owed under the lease.', TEAL, 'white', 10.5)
ax.text(.68,.20,'Important: this split applies to the commission pool, not to the full rent or deposit.', fontsize=9, color=NAVY, wrap=True)
COMM = save_fig(fig,'02_commission_waterfall.png')

# 3. Verification + escrow
fig, ax = plt.subplots(figsize=(12,6))
ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off')
ax.text(.04,.93,'3. Verification and escrow earning paths', fontsize=17, weight='bold', color=NAVY)
ax.text(.04,.88,'Two flows pay for two different kinds of value: verification work and transaction protection.', fontsize=10.5, color=MUTED)
box(ax,.06,.56,.22,.20,'OWNER / AGENT','Requests verification\nPays 1,500 XAF example', NAVY)
box(ax,.39,.56,.22,.20,'MODERATOR','Visits / confirms / uploads\nphoto + location evidence', GREEN)
box(ax,.72,.56,.22,.20,'VERIFIED LISTING','Freshness date +\nverification level shown', TEAL)
arrow(ax,.28,.66,.39,.66,'verification fee')
arrow(ax,.61,.66,.72,.66,'evidence')
box(ax,.06,.20,.22,.20,'SEEKER / OWNER','Uses protected escrow\nPays 1,500 XAF handling fee example', CORAL)
box(ax,.39,.20,.22,.20,'PLATFORM','Coordinates release,\nreceipt, dispute pathway', GOLD, INK)
box(ax,.72,.20,.22,.20,'PAYMENT PROVIDER','Processes mobile money\nand deducts its own fee', '#6B7280')
arrow(ax,.28,.30,.39,.30,'escrow fee')
arrow(ax,.61,.30,.72,.30,'settlement')
ax.text(.06,.07,'Illustrative verification split: 1,200 XAF to moderator / 300 XAF platform processing share. Actual fees require legal, provider, and operational validation.', fontsize=9.2, color=NAVY)
VER = save_fig(fig,'03_verification_escrow.png')

# 4. anti leakage loop
fig, ax = plt.subplots(figsize=(12,6))
ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off')
ax.text(.04,.93,'4. Anti-leakage loop: make the platform worth staying on', fontsize=17, weight='bold', color=NAVY)
ax.text(.04,.88,'The goal is not zero leakage; it is to make the first transaction safer and more valuable on-platform.', fontsize=10.5, color=MUTED)
box(ax,.08,.60,.22,.18,'1. MASKED CONTACT','Show neighborhood and\nproxy contact first', NAVY)
box(ax,.39,.60,.22,.18,'2. PROTECTED VISIT','Track request, time,\nagent ownership, and visit', TEAL)
box(ax,.70,.60,.22,.18,'3. ESCROW + RECEIPT','Proof of payment,\ndispute path, audit trail', GOLD, INK)
box(ax,.70,.22,.22,.18,'4. REPUTATION','Deals completed on-platform\nbuild visible trust', GREEN)
box(ax,.39,.22,.22,.18,'5. PARTNER BENEFIT','Agents keep deal proof;\nmoderators keep credibility', CORAL)
box(ax,.08,.22,.22,.18,'6. REPEAT VALUE','Renewal reminders,\nrecords, referrals, templates', NAVY)
arrow(ax,.30,.69,.39,.69)
arrow(ax,.61,.69,.70,.69)
arrow(ax,.81,.60,.81,.40)
arrow(ax,.70,.31,.61,.31)
arrow(ax,.39,.31,.30,.31)
arrow(ax,.19,.40,.19,.60)
LEAK = save_fig(fig,'04_anti_leakage_loop.png')

# 5. monthly example chart
fig, ax = plt.subplots(figsize=(10,5.2))
labels=['Agent\ncommission','Moderator\ncommission','Moderator\nverification','Platform\ncommission','Platform\nverification','Escrow\nhandling*']
vals=[140000,20000,24000,40000,6000,7500]
colors=[CORAL,GREEN,GREEN,GOLD,GOLD,TEAL]
ax.bar(labels,vals,color=colors,width=.62)
ax.set_title('Illustrative monthly economics: 5 closed deals + 20 verifications', loc='left', color=NAVY, weight='bold', fontsize=15)
ax.set_ylabel('XAF'); ax.grid(axis='y', alpha=.18); ax.spines[['top','right']].set_visible(False)
for i,v in enumerate(vals): ax.text(i,v+3000,f'{v:,}',ha='center',fontsize=9,color=INK)
ax.text(.01,-.20,'*Escrow handling shown as gross fee before payment-provider costs, refunds, compliance costs, or disputes.', transform=ax.transAxes, fontsize=8.5, color=MUTED)
MONTH = save_fig(fig,'05_monthly_economics.png')

# Create document

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill.replace('#',''))

def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color.replace('#',''))
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for i,h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], NAVY); set_cell_text(table.rows[0].cells[i], h, True, 'FFFFFF', 9)
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            set_cell_shading(cells[i], 'F7F3EC' if ridx%2==0 else 'FFFFFF'); set_cell_text(cells[i], val, False, INK, 8.8)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.style=f'Heading {level}'; return p

def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(MUTED.replace('#',''))

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].font.color.rgb=RGBColor.from_string(INK.replace('#',''))
for st in ['Heading 1','Heading 2','Heading 3']:
    styles[st].font.name='Aptos Display'; styles[st].font.color.rgb=RGBColor.from_string(NAVY.replace('#',''))
styles['Heading 1'].font.size=Pt(20); styles['Heading 2'].font.size=Pt(14); styles['Heading 3'].font.size=Pt(11)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('AFFORDABLE HOUSING CAMEROON'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(TEAL.replace('#',''))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Business Concept Plan'); r.bold=True; r.font.size=Pt(29); r.font.color.rgb=RGBColor.from_string(NAVY.replace('#',''))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Revenue sharing, participant earnings, trust operations, and anti-leakage design'); r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(MUTED.replace('#',''))

# cover callout
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('A marketplace where agents, moderators, owners, seekers, payment providers, and the platform each receive clear value from a completed housing journey.'); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(NAVY.replace('#',''))

doc.add_picture(str(ECO), width=Inches(6.9)); add_caption(doc,'Figure 1. The operating ecosystem and value exchange.')

add_heading(doc,'1. Executive concept',1)
doc.add_paragraph('Affordable Housing Cameroon should be positioned as a transaction and trust layer for rental housing, not merely as a directory of phone numbers. The platform helps seekers find relevant homes, helps owners and agents receive qualified demand, pays local moderators for credible verification work, and records the transaction so that participants receive protection and reputation value.')
doc.add_paragraph('The core commercial principle is simple: listing is free to maximize supply; the platform earns when value is created through verification, completed transactions, protected settlement, and later optional services. The figures in this document are illustrative planning assumptions and must be validated with local legal, payment-provider, and market advice before implementation.')

add_heading(doc,'2. Participants and how each person earns',1)
add_table(doc,['Participant','What they contribute','How they earn or benefit','Primary incentive'],[
['Seeker','Budget, housing need, attention, feedback','Does not receive a cash payout; earns protection, transparent costs, proof of payment, and a dispute pathway','Find a real home with less wasted time and lower fraud exposure'],
['Owner / landlord','Property, availability, lease terms','Receives rent and deposit owed under the lease; may later pay optional verification or visibility fees','Qualified demand and a clearer path to a serious renter'],
['Agent','Listing, viewing, negotiation, paperwork, deal support','Illustrative 70% of the closed-deal commission pool; repeat/referral opportunities','Tracked deal ownership, protected commission proof, reputation, and demand'],
['Local moderator','Listing verification, field evidence, freshness checks, scam flags','Illustrative verification fee share plus 10% of the closed-deal commission pool','Immediate payment for work plus recurring upside from credible listings'],
['Platform','Discovery, verification records, contact protection, escrow coordination, disputes, analytics','Illustrative 20% of closed-deal commission; verification processing share; escrow handling fee; later optional upsells','Transaction volume, trusted data, repeat use, and scalable revenue'],
['Mobile-money / payment provider','Payment processing and settlement rails','Provider-defined transaction fee; amount must be confirmed in contract','Payment volume and compliant processing'],
])

add_heading(doc,'3. Core revenue-sharing model',1)
doc.add_paragraph('The primary success-fee model applies when a seeker commits to a property and the transaction reaches the agreed trigger, such as signing and paying the first month plus deposit. The commission pool is separate from the rent and deposit. Rent and deposit belong to the contractual parties; they are not platform revenue.')
doc.add_picture(str(COMM), width=Inches(6.9)); add_caption(doc,'Figure 2. Illustrative 70/20/10 split of a 40,000 XAF commission pool.')
add_table(doc,['Money event','Illustrative amount','Agent','Moderator','Platform','Notes'],[
['Closed deal commission pool','40,000 XAF','28,000 XAF','4,000 XAF','8,000 XAF','Assumes one month of rent as commission; illustrative only'],
['Monthly rent','40,000 XAF','—','—','—','Rent remains owed under the lease, not split as platform revenue'],
['Deposit / advance','As agreed','—','—','—','Should be handled under documented lease and payment rules'],
])

add_heading(doc,'4. Verification-fee model',1)
doc.add_paragraph('Verification pays the moderator for work even when a property does not close immediately. This is important because a moderator should not be expected to perform field work for free while waiting for a future transaction.')
doc.add_picture(str(VER), width=Inches(6.9)); add_caption(doc,'Figure 3. Verification and escrow are separate earning paths.')
add_table(doc,['Fee event','Illustrative gross fee','Moderator share','Platform share','Important control'],[
['Listing verification','1,500 XAF','1,200 XAF','300 XAF','Paid for a defined verification task; does not imply legal title certification'],
['Escrow handling','1,500 XAF','—','Gross handling fee before provider costs','Must be legally and operationally validated; do not treat escrow float as interest revenue'],
])

add_heading(doc,'5. How to reduce disintermediation and leakage',1)
doc.add_paragraph('Leakage occurs when a seeker and agent meet through the platform and then complete the deal privately. The goal should not be to promise zero leakage. The goal is to make the first transaction safer, more traceable, and more valuable on-platform than off-platform.')
doc.add_picture(str(LEAK), width=Inches(6.9)); add_caption(doc,'Figure 4. The anti-leakage loop turns platform participation into ongoing value.')
add_table(doc,['Mechanism','Value created','Who benefits'],[
['Masked contact and approximate location','Prevents unrestricted scraping and protects the initial lead','Seeker, owner, agent, platform'],
['Timestamped deal ownership','Creates evidence that an agent introduced and worked the lead','Agent'],
['Escrow, receipt, and dispute path','Makes the platform the safer payment route','Seeker and owner'],
['Visible reputation and completed-deal history','Rewards participants who complete transactions on-platform','Agent and moderator'],
['Freshness and verification record','Preserves a credible listing status','Seeker and platform'],
['Repeat services','Adds renewal reminders, records, templates, and referrals','Seeker, agent, platform'],
])

add_heading(doc,'6. Illustrative monthly economics',1)
doc.add_paragraph('The following example uses five closed deals at an average 40,000 XAF monthly rent and 20 verification jobs at 1,500 XAF each. It is an illustration of mechanics, not a forecast.')
doc.add_picture(str(MONTH), width=Inches(6.7)); add_caption(doc,'Figure 5. Example gross earnings before operating expenses, provider fees, refunds, and disputes.')
add_table(doc,['Stream','Illustrative monthly volume','Agent / moderator earnings','Platform gross revenue'],[
['Closed-deal commission','5 × 40,000 XAF = 200,000 XAF pool','Agent 140,000 XAF; moderator 20,000 XAF','40,000 XAF'],
['Verification','20 × 1,500 XAF = 30,000 XAF','Moderator 24,000 XAF','6,000 XAF'],
['Escrow handling','5 × 1,500 XAF = 7,500 XAF','No direct commission assumed','7,500 XAF gross before provider costs'],
['Illustrative total','—','Agent + moderator = 184,000 XAF','53,500 XAF gross before costs'],
])

add_heading(doc,'7. Earnings examples for individuals',1)
doc.add_paragraph('If an agent closes four to six deals per month on rooms with monthly rent between 30,000 and 60,000 XAF, and the commission equals one month of rent, the agent’s 70% share would be approximately 84,000 to 252,000 XAF per month before personal costs. If a moderator verifies 15 to 20 listings per month at a 1,500 XAF fee and receives 1,200 XAF per verification, the verification-only income would be approximately 18,000 to 24,000 XAF, before any closed-deal commission slices.')
doc.add_paragraph('These are arithmetic examples derived from the assumptions supplied. They are not guaranteed earnings. Actual income depends on demand, property quality, local pricing, agent performance, verification volume, payment costs, disputes, and the final fee policy.')

add_heading(doc,'8. Revenue streams by maturity',1)
add_table(doc,['Revenue stream','Who pays','When to introduce','Platform role'],[
['Closed-deal commission share','Transaction participants according to transparent fee policy','Core launch revenue after trust and settlement are working','Take 20% illustrative share of the commission pool'],
['Verification processing share','Owner or agent','Early, once verification is defined and auditable','Coordinate payment record and verification evidence'],
['Escrow handling fee','Seeker, owner, or split according to policy','Only after legal and payment-provider validation','Provide safe settlement, receipts, and dispute process'],
['Featured listings','Owner or agent','Later, after meaningful search traffic exists','Sell visibility separately from verification'],
['Agent subscription tools','Agents','Later, after agents receive measurable value','Offer analytics, lead management, and workflow tools'],
])

doc.add_paragraph('The platform should not charge seekers to browse or charge owners merely to list during the early marketplace-building stage. Fees should be connected to real value: verification, visibility, settlement protection, or completed transactions.')

add_heading(doc,'9. Controls, risks, and safeguards',1)
add_table(doc,['Risk','Why it matters','Required safeguard'],[
['Rubber-stamped verification','A moderator may claim a visit without doing the work','Photo and location evidence, timestamp, random review, seeker feedback'],
['Agent or seeker leakage','The deal may move off-platform after the introduction','Masked contact, timestamped ownership, dispute proof, reputation, fair fees'],
['Escrow misuse','Holding deposits can create legal, financial, and consumer-protection obligations','Use a compliant payment partner; define release, refund, dispute, and reconciliation rules'],
['Unclear commission trigger','Participants may disagree about when a payout is owed','Define “closed deal” in writing and record acceptance evidence'],
['False income expectations','Illustrative numbers may be understood as promises','Label all examples as assumptions and publish actual payout rules'],
['Discriminatory or unsafe listings','Housing conditions may harm seekers','Moderation policy, reporting, removal, escalation, and local legal review'],
])

add_heading(doc,'10. Recommended payout sequence',1)
for text in [
'1. A property is submitted and remains unpublished until the verification checklist is complete.',
'2. The moderator submits evidence and receives the verification fee according to the approved payout policy.',
'3. A seeker requests contact; the platform records the lead and agent ownership where applicable.',
'4. The agent conducts the viewing and negotiation; the system records progress without exposing unnecessary personal data.',
'5. The parties reach the defined commitment trigger; the platform confirms the transaction evidence.',
'6. The commission pool is calculated and displayed before payout.',
'7. The platform releases the agent, moderator, and platform shares through the approved payment process.',
'8. The seeker and relevant participants can provide feedback, which affects reputation and future visibility.'
]:
    p=doc.add_paragraph(style='List Number'); p.add_run(text)

add_heading(doc,'11. Governance and implementation decisions still required',1)
doc.add_paragraph('Before implementation, the founder must confirm the exact commission trigger, who legally pays the commission, whether the 50–100% one-month norm applies to the selected segment, the identity and role of the landlord in each deal, and whether the platform can legally coordinate deposits or escrow through a licensed partner. The 70/20/10 allocation should be treated as a proposed starting assumption, not a final contract.')
add_table(doc,['Decision','Recommended starting position','Why it should be validated'],[
['Listing fee','Free','Maximizes early supply'],
['Verification fee','1,000–2,000 XAF range','Must reflect local travel and verification effort'],
['Commission pool','Transparent percentage or flat amount','Must be acceptable to seekers, owners, and agents'],
['Commission split','70% agent / 10% moderator / 20% platform illustrative','Must cover support, payment, disputes, and compliance'],
['Escrow','Use licensed provider or defer','Requires legal, consumer-protection, and reconciliation controls'],
])

add_heading(doc,'12. Conclusion',1)
doc.add_paragraph('The business model works only if every participant can identify a clear exchange of value. Seekers receive safer discovery and payment protection. Owners receive qualified demand. Agents receive tracked opportunities and protected commissions. Moderators receive immediate pay for real verification work and an incentive to preserve quality. Payment providers receive processing volume. The platform earns by coordinating the trusted transaction layer and only later by selling optional visibility or software tools.')
doc.add_paragraph('The strategic principle is therefore: **build the trust and settlement layer, not just the listing page**. Revenue sharing should be transparent before the first deal, every payout should have a recorded trigger, and every earning estimate should be presented as an assumption until validated through local operations.')

add_heading(doc,'Appendix A — Source assumptions',1)
doc.add_paragraph('This document is based on the supplied content. The commission split, verification fee, earnings examples, and scale examples are illustrative assumptions from that content and should be validated through local interviews, pilot transactions, payment-provider terms, and professional legal/accounting review. No external market statistic has been presented as a verified fact in this document.')

# Footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run('Affordable Housing Cameroon — Business Concept Plan | Illustrative model for validation'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(MUTED.replace('#',''))

out=ROOT/'affordable-housing-cameroon-business-concept-plan.docx'
doc.save(out)
print(out)
print('Assets:', ASSETS)
